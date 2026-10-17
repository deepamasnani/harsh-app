from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel
from sqlalchemy import or_
from sqlalchemy.orm import Session

from ..auth import require_admin
from ..database import get_db
from ..models import Product, User

router = APIRouter(prefix="/products", tags=["products"])


class ProductIn(BaseModel):
    sku: str
    name: str
    unit: str = "pcs"
    price: float = 0
    stock: float = 0
    category: str = "General"


def product_out(p: Product) -> dict:
    return {
        "id": p.id,
        "sku": p.sku,
        "name": p.name,
        "unit": p.unit,
        "price": p.price,
        "stock": p.stock,
        "category": p.category,
        "is_active": p.is_active,
    }


@router.get("")
def list_products(
    q: str = "",
    include_inactive: bool = False,
    db: Session = Depends(get_db),
    user: User = Depends(require_admin),
):
    query = db.query(Product)
    if not include_inactive:
        query = query.filter(Product.is_active.is_(True))
    if q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(or_(Product.sku.ilike(like), Product.name.ilike(like), Product.category.ilike(like)))
    rows = query.order_by(Product.name.asc()).all()
    return [product_out(p) for p in rows]


@router.post("")
def create_product(body: ProductIn, db: Session = Depends(get_db), user: User = Depends(require_admin)):
    sku = body.sku.strip().upper()
    if db.query(Product).filter(Product.sku == sku).first():
        raise HTTPException(status_code=400, detail="SKU already exists")
    p = Product(
        sku=sku,
        name=body.name.strip(),
        unit=body.unit.strip() or "pcs",
        price=body.price,
        stock=body.stock,
        category=body.category.strip() or "General",
    )
    db.add(p)
    db.commit()
    db.refresh(p)
    return product_out(p)


@router.put("/{product_id}")
def update_product(
    product_id: int,
    body: ProductIn,
    db: Session = Depends(get_db),
    user: User = Depends(require_admin),
):
    p = db.get(Product, product_id)
    if not p:
        raise HTTPException(status_code=404, detail="Product not found")
    sku = body.sku.strip().upper()
    clash = db.query(Product).filter(Product.sku == sku, Product.id != product_id).first()
    if clash:
        raise HTTPException(status_code=400, detail="SKU already exists")
    p.sku = sku
    p.name = body.name.strip()
    p.unit = body.unit.strip() or "pcs"
    p.price = body.price
    p.stock = body.stock
    p.category = body.category.strip() or "General"
    p.is_active = True
    db.commit()
    return product_out(p)


@router.delete("/{product_id}")
def delete_product(product_id: int, db: Session = Depends(get_db), user: User = Depends(require_admin)):
    p = db.get(Product, product_id)
    if not p:
        raise HTTPException(status_code=404, detail="Product not found")
    p.is_active = False
    db.commit()
    return {"ok": True}


def _cell(row: dict, *keys, default=""):
    for k in keys:
        for rk, rv in row.items():
            if str(rk).strip().lower() == k:
                return rv if rv is not None else default
    return default


def _num(val, default=0.0) -> float:
    if val is None or val == "":
        return default
    try:
        return float(str(val).replace(",", "").strip())
    except ValueError:
        return default


@router.post("/import")
async def import_products(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(require_admin),
):
    name = (file.filename or "").lower()
    data = await file.read()
    rows = []

    if name.endswith(".csv"):
        import csv
        import io

        text = data.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
    elif name.endswith(".xlsx") or name.endswith(".xls"):
        import io

        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), data_only=True)
        ws = wb.active
        headers = [str(c.value).strip() if c.value is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for excel_row in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None or str(v).strip() == "" for v in excel_row):
                continue
            rows.append({headers[i]: excel_row[i] if i < len(excel_row) else None for i in range(len(headers))})
    else:
        raise HTTPException(status_code=400, detail="Upload a .xlsx or .csv file")

    created = updated = skipped = 0
    for row in rows:
        sku = str(_cell(row, "sku", "item code", "code", "item_code") or "").strip().upper()
        pname = str(_cell(row, "name", "item", "product", "description") or "").strip()
        if not sku or not pname:
            skipped += 1
            continue
        unit = str(_cell(row, "unit", "uom") or "pcs").strip() or "pcs"
        price = _num(_cell(row, "price", "rate", "mrp", "selling_price"))
        stock = _num(_cell(row, "stock", "qty", "quantity", "qty_on_hand"))
        category = str(_cell(row, "category", "group") or "General").strip() or "General"

        existing = db.query(Product).filter(Product.sku == sku).first()
        if existing:
            existing.name = pname
            existing.unit = unit
            existing.price = price
            existing.stock = stock
            existing.category = category
            existing.is_active = True
            updated += 1
        else:
            db.add(
                Product(sku=sku, name=pname, unit=unit, price=price, stock=stock, category=category)
            )
            created += 1

    db.commit()
    return {"created": created, "updated": updated, "skipped": skipped, "total": len(rows)}
